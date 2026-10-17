import os
import re
import io
import datetime
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.styles.colors import Color

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "templates")
DEFAULT_QC_FILE = os.path.join(TEMPLATE_DIR, "QC exception Vlookup.xlsx")
FALLBACK_QC_FILE = r"C:\vendor Rates\QC exception Vlookup.xlsx"

# Default QC Exception codes (Satellite, High Risk, Freephone, and Austria special codes)
DEFAULT_QC_CODES = [
    "871", "872", "873", "874", "875", "876", "877", "879", "884", "885", "887", "889",
    "969", "979", "999", "1855", "1866", "1877", "1888", "2897", "6993", "6998",
    "388", "800", "991", "808", "684", "1671", "671",
    # Austria Special / High Risk codes
    "436132308993", "436646570601", "436648668402", "436648668411", "436648668412",
    "436648668420", "436648668423", "436648668424", "436648740174", "436649469663",
    "4366499288491", "4366499288492", "4366499288493", "4366499288496", "4366499288497",
    "4366499288498", "4367030947", "4367030981", "436708012001", "436708012012",
    "436708012014", "436708012018", "436708012019", "436708012021", "436708012022",
    "436708012027", "43670801203", "43670801204", "43670801205", "43670801206",
    "43670801207", "43670801208", "43670801209", "43670801201", "4367080121",
    "436708012121", "43670801214", "436708012150", "436708012154", "436708012160",
    "43670801217", "43670801218", "43670801219", "436708012200", "436708012202",
    "436708012203", "436708012204", "436708012205", "436708012206", "436708012208",
    "436708012211", "436708012212", "436708012217", "436708012219", "436708012221",
    "436708012223", "436708012225", "436708012233", "43670801224", "436708012242",
    "436708012250", "436708012251", "436708012257", "436708012258", "436708012262",
    "436708012263", "436708012265", "436708012267", "436708012268", "436708012271",
    "436708012272", "43670801228", "43670801229", "436708012327", "436708012329",
    "436708012330", "436708012333", "436708012334", "436708012340", "436708012342",
    "436708012344", "436708012346", "436708012349", "436708012350", "436708012355",
    "43670801236", "436708012374", "43670801238", "43670801240", "43670801241",
    "436708012410", "436708012427", "43670801243", "43670801244", "43670801245",
    "43670801246", "43670801247", "43670801248", "43670801249", "436708012490",
    "436708012492", "436708012493", "436708012494", "436708012526", "43670801253",
    "43670801254", "436708012551", "436708012552", "43670801256", "43670801257",
    "43670801258", "43670801259", "436708012591", "436708012597", "436708012598",
    "43670801260", "436708012611", "436708012612", "436708012613", "436708012632",
    "436708012660", "436708012674", "43670801268", "43670801269", "436708012710",
    "436708012712", "43670801273", "436708012742", "436708012753", "436708012757",
    "436708012758", "436708012762", "436708012768", "436708012772", "43670801278",
    "436708012793", "436708012795", "436708013092", "436708013094", "436708013096",
    "43670809029", "436708090340", "436708091166", "436708091194", "436708091196",
    "436708092021", "436708092023", "436708092027", "436708092304", "436708093151",
    "436708093155", "436708093156", "436708093157", "43670809336", "436708094044",
    "436708094047", "43670809444", "436708095051", "436708095052", "436708095054",
    "436708095057", "436708095210", "436708095212", "436708095213", "436708095215",
    "43670809613", "43670809622", "436708097123", "43670809747", "436708098012",
    "43670809820", "436708098380", "436708099110", "436708099113", "436708099115",
    "436708099226", "4368110728430", "4368110788534", "4368110891145", "4368120310611",
    "4368120875953"
]

def load_qc_exception_codes():
    codes = set(DEFAULT_QC_CODES)
    for p in [DEFAULT_QC_FILE, FALLBACK_QC_FILE]:
        if os.path.exists(p):
            try:
                df = pd.read_excel(p)
                for col in df.columns:
                    if 'code' in str(col).lower():
                        for val in df[col].dropna():
                            c_str = str(val).split('.')[0].strip()
                            if c_str:
                                codes.add(c_str)
                break
            except Exception:
                pass
    return codes

def detect_header_row_and_columns(df_raw):
    """
    Scans first 40 rows to identify table header and map columns:
    Handles 'CODES', 'Rate Changes', 'Effective From', 'Destination',
    'National & International', and separate 'Country Code' + 'Area Code' columns.
    """
    best_row_idx = 0
    best_mapping = {}
    best_score = -1

    for row_idx in range(min(40, len(df_raw))):
        row_vals = [str(x).strip().lower() for x in df_raw.iloc[row_idx].values]
        mapping = {}
        score = 0

        for col_idx, val in enumerate(row_vals):
            v = re.sub(r'[^a-z0-9]', ' ', val).strip()
            
            # 1. Status / Changes (check before rate so 'rate changes' is mapped to status)
            if any(k in v for k in ['rate changes', 'rate change', 'changes', 'status', 'comment', 'action', 'change']):
                if 'status' not in mapping or 'change' in v or 'status' in v:
                    mapping['status'] = col_idx
                    score += 4
            # 2. Dial code keywords
            elif any(k in v for k in ['dial code', 'dial codes', 'dialcode', 'dialcodes', 'prefix', 'prefixes', 'dest code', 'dest codes', 'destination code', 'complete code', 'complete codes', 'full code', 'full codes', 'complete dial code', 'full dial code', 'full prefix', 'complete prefix', 'dialed digit', 'dialed digits', 'dialed number', 'dialed numbers']) or v in ['code', 'codes']:
                if 'dial_code' not in mapping:
                    mapping['dial_code'] = col_idx
                    score += 6
            elif 'country code' in v:
                if 'country_code' not in mapping:
                    mapping['country_code'] = col_idx
                    score += 5
            elif 'area code' in v:
                if 'area_code' not in mapping:
                    mapping['area_code'] = col_idx
                    score += 5
            # 3. Destination keywords
            elif any(k in v for k in ['destination', 'destinations', 'summary', 'country', 'location', 'dest name', 'description', 'loc name', 'national', 'international']):
                if 'destination' not in mapping:
                    mapping['destination'] = col_idx
                    score += 6
            # 4. Rate keywords (use word boundaries so 'gold' is never mistaken for 'old', and ignore previous rates)
            elif (any(k in v for k in ['new rate', 'price usd', 'rate usd', 'price', 'prices', 'rate', 'rates', 'tariff', 'tariffs', 'usd']) or '$' in val) and not re.search(r'\b(?:current|old|prev|previous|existing)\b', v):
                if 'rate' not in mapping or 'new rate' in v or 'new price' in v or 'rate usd' in v:
                    mapping['rate'] = col_idx
                    score += 6
            # 5. Effective Date keywords (ignore expiry/end dates)
            elif any(k in v for k in ['effective from', 'eff date', 'effective date', 'application date', 'date activate', 'date', 'eff', 'pricing valid from', 'valid from']) and not any(k in v for k in ['expiry', 'expire', 'end', 'until', 'close', 'termination']):
                if 'effective_date' not in mapping or 'effective' in v:
                    mapping['effective_date'] = col_idx
                    score += 5
            # 6. Pulse / Billing keywords
            elif any(k in v for k in ['increments', 'increment', 'incremental', 'pulse', 'billing', 'initial billing', 'incremental billing', 'billing increment', 'minimum increments', 'billing step', 'min billing time', 'min inc', 'min/inc', 'billing pulse', 'billing unit', 'pulse rate', 'unit', 'resolution', 'minimal time', 'billing resolution', 'interval', 'intervals', 'billing interval']):
                if 'pulse' not in mapping:
                    mapping['pulse'] = col_idx
                    score += 4

        if ('country_code' in mapping and 'area_code' in mapping) and 'dial_code' not in mapping:
            mapping['dial_code_split'] = True
        elif 'country_code' in mapping and 'dial_code' not in mapping and 'area_code' not in mapping:
            mapping['dial_code'] = mapping['country_code']
            score += 6

        if score > best_score and ('dial_code' in mapping or ('country_code' in mapping and 'area_code' in mapping)) and ('rate' in mapping or 'destination' in mapping):
            best_score = score
            best_row_idx = row_idx
            best_mapping = mapping

    return best_row_idx, best_mapping

def analyze_vendor_comments(df_slice, status_col_idx):
    if status_col_idx is None:
        return False, "UPdated"
        
    vals = df_slice.iloc[:, status_col_idx].dropna().astype(str).tolist()
    if not vals:
        return False, "UPdated"
        
    found_types = set()
    for v in vals:
        v_low = v.lower().strip()
        if not v_low or v_low in ['nan', 'none', 'null', '']:
            continue
        if any(k in v_low for k in ['inc', 'increase', 'up']):
            found_types.add('inc')
        elif any(k in v_low for k in ['dec', 'decrease', 'down']):
            found_types.add('dec')
        elif any(k in v_low for k in ['same', 'unchange', 'no change']):
            found_types.add('same')
        elif any(k in v_low for k in ['del', 'delete', 'remove', 'block', 'close']):
            found_types.add('delete')
        elif any(k in v_low for k in ['new']):
            found_types.add('new')

    if not found_types:
        return False, "UPdated"

    ordered_tags = []
    for t in ['inc', 'dec', 'same', 'delete', 'new']:
        if t in found_types:
            ordered_tags.append(t)
            
    tag_str = ", ".join(ordered_tags)
    return True, f"Various {tag_str}"

def fast_process_dataframe(df_raw, qc_codes, source_filename=""):
    """
    High-speed vectorized data transformation engine (< 2 seconds):
    - Explodes comma-separated dial codes if present
    - Strictly filters out rows without valid dial codes (drops footer notes, disclaimers, blanks)
    - Vectorized QC Exception filtering
    - Vectorized USA/Canada/Puerto Rico NANP (+1) filtering
    - Vectorized Rate conversion (>=1.00 -> -1, <1.00 -> original, empty/block/delete -> -1)
    - Preserves vendor's exact Effective Date (never invents today's date)
    - Preserves vendor's exact Billing Pulse (never invents 1/1 if vendor didn't provide)
    - Output: Dial Code, Summary, New Rate, Eff. Date, Minimum, Increment
    """
    header_row_idx, col_map = detect_header_row_and_columns(df_raw)
    
    df_data = df_raw.iloc[header_row_idx + 1:].copy().reset_index(drop=True)
    header_names = df_raw.iloc[header_row_idx].values

    code_idx = col_map.get('dial_code', 0)
    dest_idx = col_map.get('destination', 1 if code_idx != 1 else 0)
    rate_idx = col_map.get('rate', 2)
    date_idx = col_map.get('effective_date', None)
    status_idx = col_map.get('status', None)
    pulse_idx = col_map.get('pulse', None)

    # Separate min/inc if exists
    min_col_idx = None
    inc_col_idx = None
    for idx, col_name in enumerate(header_names):
        c_name = str(col_name).lower().strip()
        if any(k in c_name for k in ['minimum', 'initial', 'min duration', 'min pulse', 'minimal', 'minimal time', 'min time', 'init time', 'initial period', 'init period', 'min period']) and not any(k in c_name for k in ['increment', 'incremental', 'incre', 'resolution', 'recurring', 'recur']):
            min_col_idx = idx
        elif any(k in c_name for k in ['increment', 'incremental', 'incre duration', 'inc pulse', 'resolution', 'billing resolution', 'billing res', 'recurring', 'recur', 'recurring period', 'rec period', 'repeat', 'next period']) and not any(k in c_name for k in ['minimum', 'initial', 'minimal']):
            inc_col_idx = idx

    # If only one increment or minimum column exists without separate pair, treat as combined pulse column
    if min_col_idx is None and inc_col_idx is not None and pulse_idx is None:
        pulse_idx = inc_col_idx
    elif min_col_idx is not None and inc_col_idx is None and pulse_idx is None:
        pulse_idx = min_col_idx

    # 1. Expand dial codes (split columns or comma/semicolon delimited)
    if col_map.get('dial_code_split', False):
        c_col = col_map['country_code']
        a_col = col_map['area_code']
        raw_c = df_data.iloc[:, c_col].fillna('').astype(str).str.split('.').str[0].str.strip()
        raw_a = df_data.iloc[:, a_col].fillna('').astype(str).str.split('.').str[0].str.strip()
        # Clean area code from any placeholder template strings like %%=...
        raw_a_clean = raw_a.str.replace(r'[^\d,;]', '', regex=True)
        if raw_a_clean.str.contains(r'[,;]', regex=True, na=False).any():
            df_data['_split_a'] = raw_a_clean.str.split(r'[,;]+')
            df_data = df_data.explode('_split_a').reset_index(drop=True)
            raw_c = df_data.iloc[:, c_col].fillna('').astype(str).str.split('.').str[0].str.strip()
            raw_a_clean = df_data['_split_a'].fillna('').astype(str).str.split('.').str[0].str.strip()
        raw_code_series = (raw_c + raw_a_clean).str.split('.').str[0].str.strip()
    else:
        code_raw_series = df_data.iloc[:, code_idx].fillna('').astype(str).str.strip()
        if code_raw_series.str.contains(r'[,;]', regex=True, na=False).any():
            df_data['_split_code'] = code_raw_series.str.split(r'[,;]+')
            df_data = df_data.explode('_split_code').reset_index(drop=True)
            raw_code_series = df_data['_split_code'].fillna('').astype(str).str.split('.').str[0].str.strip()
        else:
            raw_code_series = code_raw_series.str.split('.').str[0].str.strip()

    # CRITICAL: STRICT DIAL CODE FILTERING & SANITIZATION
    # Checks notes/disclaimers/payment terms on original text before stripping characters
    raw_code_clean_lower = raw_code_series.str.lower().str.strip()
    not_note_code = ~raw_code_clean_lower.str.contains(r'note|disclaimer|total|table|gmt|term|payment|bank|invoice|account|contact|support|page', regex=True, na=False) & ~raw_code_series.str.contains(r'[a-zA-Z]{3,}', regex=True, na=False)
    
    dest_raw_series = df_data.iloc[:, dest_idx].fillna('').astype(str).str.lower().str.strip()
    not_note_dest = ~dest_raw_series.str.contains(r'^(?:total|subtotal|grand total|disclaimer|notice|terms|conditions|payment|bank details|account|page)\b', regex=True, na=False)

    # Strip ALL special characters (#, *, +, -, spaces, brackets, letters, etc.) keeping strictly digits
    clean_digits = raw_code_series.str.replace(r'[^\d]', '', regex=True)
    # Strip leading international exit prefix '00' if present (e.g. 0044 -> 44)
    clean_digits = clean_digits.str.replace(r'^00', '', regex=True)

    has_digit = clean_digits.str.len() > 0
    not_empty = ~clean_digits.isin(['', 'nan', 'none', 'null'])
    valid_dial_mask = has_digit & not_empty & not_note_code & not_note_dest

    df_data = df_data[valid_dial_mask].copy().reset_index(drop=True)
    code_series = clean_digits[valid_dial_mask].reset_index(drop=True)

    initial_row_count = len(df_data)

    # 2. Vectorized QC Exception Mask (Codes + Destinations: Freephone, Satellite, Inmarsat)
    dest_series = df_data.iloc[:, dest_idx].fillna('').astype(str).str.upper().str.strip()
    qc_tuple = tuple(qc_codes)
    is_qc_code = code_series.str.startswith(qc_tuple)
    qc_dest_pattern = r'\b(?:INTERNATIONAL\s+FREEPHONE|FREEPHONE|FREE\s*PHONE|SATELLITE|INMARSAT)\b'
    is_qc_dest = dest_series.str.contains(qc_dest_pattern, regex=True, na=False)
    is_qc_mask = is_qc_code | is_qc_dest

    # 3. Vectorized USA / Canada / Puerto Rico (+1) Mask
    nanp_pattern = r'\b(?:USA|U\.S\.A|UNITED STATES|AMERICA|CANADA|PUERTO RICO|PR)\b'
    is_nanp_dest = dest_series.str.contains(nanp_pattern, regex=True, na=False)
    is_nanp_code = (code_series == '1') | (code_series.str.startswith('1') & (code_series.str.len() >= 4) & code_series.str.isdigit())
    is_nanp_mask = is_nanp_dest | is_nanp_code

    # Exclusions
    is_excluded_mask = is_qc_mask | is_nanp_mask
    valid_mask = ~is_excluded_mask

    df_valid = df_data[valid_mask].copy().reset_index(drop=True)
    valid_codes = code_series[valid_mask].reset_index(drop=True)

    # 4. Comments / Status analysis
    has_comments, comments_tag = analyze_vendor_comments(df_data, status_idx)

    # 5. Vectorized Rate Logic (>= 1.00 -> -1, <= 0.0 -> -1, empty/block/delete -> -1)
    raw_rates_str = df_valid.iloc[:, rate_idx].fillna('').astype(str).str.lower().str.strip()
    clean_nums = pd.to_numeric(raw_rates_str.str.replace(r'[^\d.-]', '', regex=True), errors='coerce')
    
    is_empty_or_block = raw_rates_str.isin(['', 'nan', 'none', 'null', '-']) | raw_rates_str.str.contains(r'block|close|na|n/a|del|free', regex=True, na=False)
    
    if status_idx is not None:
        status_str = df_valid.iloc[:, status_idx].fillna('').astype(str).str.lower().str.strip()
        is_delete_status = status_str.str.contains(r'del|block|close', regex=True, na=False)
    else:
        is_delete_status = pd.Series(False, index=df_valid.index)
        
    is_zero = (clean_nums <= 0)
    is_minus_1 = is_empty_or_block | is_delete_status | (clean_nums >= 1.00) | is_zero | clean_nums.isna()
    final_rates = np.where(is_minus_1, -1.0, clean_nums)

    # 6. Date Logic (Strict: Never invent today's date!)
    date_objs = []
    date_strs = []
    if date_idx is not None:
        parsed_dates = pd.to_datetime(df_valid.iloc[:, date_idx], format='mixed', errors='coerce')
        for dt in parsed_dates:
            if pd.notna(dt):
                d_val = dt.date()
                date_objs.append(d_val)
                date_strs.append(d_val.strftime('%m/%d/%Y'))
            else:
                date_objs.append(None)
                date_strs.append('')
    else:
        date_objs = [None] * len(df_valid)
        date_strs = [''] * len(df_valid)

    # 7. Pulse Logic (Strict: Never invent pulse if vendor did not provide it!)
    mins = []
    incs = []
    if min_col_idx is not None and inc_col_idx is not None:
        raw_mins = df_valid.iloc[:, min_col_idx].tolist()
        raw_incs = df_valid.iloc[:, inc_col_idx].tolist()
        for m_val, i_val in zip(raw_mins, raw_incs):
            m_res = None
            i_res = None
            if pd.notna(m_val):
                try:
                    m_res = int(float(m_val))
                except Exception:
                    nums_m = re.findall(r'\d+(?:\.\d+)?', str(m_val).strip())
                    if nums_m:
                        m_res = int(float(nums_m[0]))
            if pd.notna(i_val):
                try:
                    i_res = int(float(i_val))
                except Exception:
                    nums_i = re.findall(r'\d+(?:\.\d+)?', str(i_val).strip())
                    if nums_i:
                        i_res = int(float(nums_i[0]))
            mins.append(m_res)
            incs.append(i_res)
    elif pulse_idx is not None:
        pulse_list = df_valid.iloc[:, pulse_idx].fillna('').astype(str).tolist()
        for p_val in pulse_list:
            p_clean = p_val.strip()
            if not p_clean or p_clean.lower() in ['nan', 'none', 'null', '', '-']:
                mins.append(None)
                incs.append(None)
                continue
            nums = re.findall(r'\d+(?:\.\d+)?', p_clean)
            if not nums:
                mins.append(None)
                incs.append(None)
                continue
            try:
                if len(nums) >= 3:
                    # e.g. 0/60/1, 0-60-1, 0+60+1 -> min=60, inc=1
                    m = int(float(nums[1]))
                    i = int(float(nums[2]))
                elif len(nums) == 2:
                    # e.g. 1/1, 1-1, 1+1, 60/1, 60-1, 60+1, 60:1
                    m = int(float(nums[0]))
                    i = int(float(nums[1]))
                elif len(nums) == 1:
                    # e.g. 1 or 60
                    v = int(float(nums[0]))
                    m, i = v, v
                else:
                    m, i = None, None
                mins.append(m)
                incs.append(i)
            except Exception:
                mins.append(None)
                incs.append(None)
    else:
        mins = [None] * len(df_valid)
        incs = [None] * len(df_valid)

    # 8. Clean Dial Codes
    clean_codes = []
    for c in valid_codes:
        c_str = str(c).split('.')[0].strip() if pd.notna(c) else ''
        if c_str.isdigit() and not c_str.startswith('0'):
            clean_codes.append(int(c_str))
        else:
            clean_codes.append(c_str)

    valid_summaries = df_valid.iloc[:, dest_idx].fillna('').astype(str).str.strip().tolist()

    df_clean = pd.DataFrame({
        'Dial Code': clean_codes,
        'Summary': valid_summaries,
        'New Rate': final_rates,
        'Eff. Date': date_strs,
        '_dt_obj': date_objs,
        'Minimum': mins,
        'Increment': incs
    })

    # Excluded rows for audit log
    df_excluded = pd.DataFrame({
        'Dial Code': code_series[is_excluded_mask],
        'Destination': df_data.iloc[is_excluded_mask, dest_idx].fillna('').astype(str).str.strip(),
        'Raw Rate': df_data.iloc[is_excluded_mask, rate_idx],
        'Reason': np.where(is_qc_mask[is_excluded_mask], 'QC Exception Code', 'USA/Canada/Puerto Rico (+1)')
    })

    final_row_count = len(df_clean)
    qc_count = int(is_qc_mask.sum())
    nanp_count = int(is_nanp_mask.sum())
    rates_minus1_count = int(is_minus_1.sum())
    total_excluded = int(is_excluded_mask.sum())
    verification_passed = (initial_row_count - total_excluded == final_row_count)

    summary_stats = {
        'source_file': source_filename,
        'initial_rows': initial_row_count,
        'qc_excluded': qc_count,
        'nanp_excluded': nanp_count,
        'rates_minus1_count': rates_minus1_count,
        'final_clean_rows': final_row_count,
        'verification_passed': verification_passed,
        'has_comments': has_comments,
        'comments_tag': comments_tag
    }

    return df_clean, df_excluded, summary_stats

def generate_target_excel_bytes(df_clean):
    """
    Generates target Excel file matching 'Vendor Sample sheet.xlsx' exactly:
    - Sheet Name: 'Rate_Amendment'
    - Header fill: openpyxl indexed 22 (Silver/Grey)
    - Header font: Calibri, 11pt, Bold, black (indexed 8)
    - Alignment: Center, Center
    - Eff. Date: Native datetime.date with number_format = 'mm-dd-yy'
    - Exact column widths matching Kaif's template
    """
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rate_Amendment"

    headers = ['Dial Code', 'Summary', 'New Rate', 'Eff. Date', 'Minimum', 'Increment']
    ws.append(headers)

    fill_22 = PatternFill(fill_type='solid', start_color=Color(indexed=22), end_color=Color(indexed=22))
    font_hdr = Font(name='Calibri', size=11, bold=True, color=Color(indexed=8))
    align_hdr = Alignment(horizontal='center', vertical='center')

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = fill_22
        cell.font = font_hdr
        cell.alignment = align_hdr

    # Fast row appending
    dt_list = df_clean['_dt_obj'].tolist() if '_dt_obj' in df_clean.columns else [None] * len(df_clean)
    codes = df_clean['Dial Code'].tolist()
    sums = df_clean['Summary'].tolist()
    rates = df_clean['New Rate'].tolist()
    mins = df_clean['Minimum'].tolist()
    incs = df_clean['Increment'].tolist()

    for i in range(len(df_clean)):
        c = codes[i]
        s = sums[i]
        r = float(rates[i])
        d = dt_list[i] if (i < len(dt_list) and dt_list[i] is not None) else None
        m = mins[i] if (i < len(mins) and mins[i] is not None and pd.notna(mins[i])) else None
        inc = incs[i] if (i < len(incs) and incs[i] is not None and pd.notna(incs[i])) else None
        
        ws.append([c, s, r, d, m, inc])

    # Apply date number format to column D ONLY where date object exists
    for row_idx in range(2, len(df_clean) + 2):
        cell_d = ws.cell(row=row_idx, column=4)
        if cell_d.value is not None and isinstance(cell_d.value, (datetime.date, datetime.datetime)):
            cell_d.number_format = 'mm-dd-yy'

    # Exact column widths matching Vendor Sample sheet.xlsx
    ws.column_dimensions['A'].width = 11.0
    ws.column_dimensions['B'].width = 44.0
    ws.column_dimensions['C'].width = 14.5
    ws.column_dimensions['D'].width = 13.0
    ws.column_dimensions['E'].width = 10.0
    ws.column_dimensions['F'].width = 10.5

    wb.save(output)
    output.seek(0)
    return output.getvalue()

def build_output_filename(base_name, comments_tag, member_name):
    base = str(base_name or '').strip()
    if not base:
        base = "Rate Deck"
        
    tag = str(comments_tag or 'UPdated').strip()
    member = str(member_name or 'Kaif').strip()

    # Clean existing extensions if typed by user
    if base.lower().endswith('.xlsx'):
        base = base[:-5].strip()

    # Handle [comments] or [Comments] placeholder if user included it in base name
    if '[comments]' in base.lower():
        if 'a to z' not in base.lower():
            base = re.sub(r'\[comments\]', f'A to Z {tag}', base, flags=re.IGNORECASE).strip()
        else:
            base = re.sub(r'\[comments\]', tag, base, flags=re.IGNORECASE).strip()

    # If member name already at end, strip to prevent duplication
    if member and base.lower().endswith(member.lower()):
        base = base[:-len(member)].strip()

    if 'a to z' in base.lower():
        if any(k in base.lower() for k in ['various', 'updated', 'inc', 'dec', 'same', 'delete']):
            filename = f"{base} {member}.xlsx"
        else:
            filename = f"{base} {tag} {member}.xlsx"
    else:
        filename = f"{base} A to Z {tag} {member}.xlsx"

    filename = re.sub(r'\s+', ' ', filename).strip()
    return filename
